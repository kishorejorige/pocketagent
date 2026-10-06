import os
import docx
from unittest.mock import MagicMock

os.environ.setdefault("GEMINI_API_KEY", "test")

from app.memory import rag


def test_chunking():
    text = "\n\n".join(f"Paragraph {i} " + "x" * 300 for i in range(10))
    chunks = rag.chunk_text(text, size=900)
    assert len(chunks) > 1
    assert all(len(c) <= 900 for c in chunks)


def test_cosine():
    assert abs(rag._cos([1, 0], [1, 0]) - 1) < 1e-9
    assert rag._cos([1, 0], [0, 1]) == 0


def test_load_text_txt_and_md(tmp_path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Hello plain text", encoding="utf-8")
    md_file = tmp_path / "sample.md"
    md_file.write_text("# Hello markdown", encoding="utf-8")

    assert rag.load_text(txt_file) == "Hello plain text"
    assert rag.load_text(md_file) == "# Hello markdown"


def test_load_text_docx(tmp_path):
    doc_path = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph text.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Value 1"
    doc.save(str(doc_path))

    extracted = rag.load_text(doc_path)
    assert "First paragraph text." in extracted
    assert "Header 1 | Value 1" in extracted


def test_load_text_pdf(tmp_path, monkeypatch):
    pdf_path = tmp_path / "test.pdf"
    pdf_path.write_bytes(b"%PDF-fake")

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "PDF Page 1 Content"
    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1]

    monkeypatch.setattr("pypdf.PdfReader", lambda path: mock_reader)
    assert rag.load_text(pdf_path) == "PDF Page 1 Content"

    # Test empty / scanned PDF
    mock_page1.extract_text.return_value = ""
    assert "no text found, it may be a scanned PDF" in rag.load_text(pdf_path)


def test_index_rebuild_and_corrupt_file_survival(tmp_path, monkeypatch):
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    index_file = tmp_path / "rag_index.json"

    monkeypatch.setattr(rag, "DOCS_DIR", docs_dir)
    monkeypatch.setattr(rag, "INDEX_FILE", index_file)
    monkeypatch.setattr(rag, "embed", lambda texts, prefix: [[1.0, 0.0] for _ in texts])

    good_file = docs_dir / "good.txt"
    good_file.write_text("Valid text content", encoding="utf-8")

    corrupt_file = docs_dir / "corrupt.pdf"
    corrupt_file.write_bytes(b"corrupt non-pdf data")

    # Rebuild index should succeed for good_file and survive corrupt_file failure
    res = rag.build_index()
    assert "1 documents" in res

    # Verify search works over good_file
    results = rag.search("Valid")
    assert len(results) == 1
    assert results[0][1] == "good.txt"

    # Test file deletion update
    good_file.unlink()
    res2 = rag.build_index()
    assert "0 documents" in res2