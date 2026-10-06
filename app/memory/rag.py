"""Local RAG: chunk documents, embed them with Ollama, search by cosine similarity.
Run `python -m app.memory.rag` to (re)build the index by hand."""
import json
import math
import urllib.request
from pathlib import Path

from app import config
from app.config import DATA_DIR

DOCS_DIR = DATA_DIR / "docs"
INDEX_FILE = DATA_DIR / "rag_index.json"
INDEX_VERSION = 2
EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def load_text(path: Path) -> str:
    """Load text content from a .txt, .md, .pdf, or .docx file."""
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    elif ext == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(str(path))
        text_parts = [page.extract_text() or "" for page in reader.pages]
        full_text = "\n\n".join(t.strip() for t in text_parts if t.strip())
        if not full_text:
            return "no text found, it may be a scanned PDF; OCR is not supported"
        return full_text
    elif ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")


def embed(texts: list[str], prefix: str) -> list[list[float]]:
    """Get embeddings from Ollama. nomic-embed-text expects a task prefix."""
    body = json.dumps(
        {"model": config.EMBED_MODEL, "input": [prefix + t for t in texts]}
    ).encode()
    req = urllib.request.Request(
        f"{config.OLLAMA_URL}/api/embed",
        data=body,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.load(r)["embeddings"]


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    """Pack paragraphs into chunks of at most `size` characters."""
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in paras:
        if len(cur) + len(p) + 2 <= size:
            cur = f"{cur}\n\n{p}" if cur else p
            continue
        if cur:
            chunks.append(cur)
        while len(p) > size:  # a single very long paragraph
            chunks.append(p[:size])
            p = p[size - overlap:]
        cur = p
    if cur:
        chunks.append(cur)
    return chunks


def _cos(a: list[float], b: list[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(x * x for x in b))
    return dot / (na * nb) if na and nb else 0.0


def _load() -> dict:
    if INDEX_FILE.exists():
        try:
            data = json.loads(INDEX_FILE.read_text(encoding="utf-8"))
            if isinstance(data, dict) and data.get("_version") == INDEX_VERSION:
                return data.get("documents", {})
        except Exception:
            pass
    return {}


def build_index() -> str:
    """Index new or changed files and drop deleted ones. Cheap when nothing changed."""
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    index = _load()
    files = {
        f.name: f for f in DOCS_DIR.iterdir()
        if f.is_file() and f.suffix.lower() in EXTENSIONS
    }
    changed = 0
    for name, path in files.items():
        try:
            mtime = path.stat().st_mtime
            if index.get(name, {}).get("mtime") == mtime:
                continue
            raw_text = load_text(path)
            chunks = chunk_text(raw_text)
            vecs: list[list[float]] = []
            for i in range(0, len(chunks), 16):
                vecs += embed(chunks[i:i + 16], "search_document: ")
            index[name] = {
                "mtime": mtime,
                "chunks": [
                    {"text": t, "vec": [round(x, 5) for x in v]}
                    for t, v in zip(chunks, vecs)
                ],
            }
            changed += 1
        except Exception as e:
            print(f"Warning: Failed to index file {name}: {e}")

    for name in list(index):
        if name not in files:
            del index[name]
            changed += 1
    if changed:
        INDEX_FILE.write_text(
            json.dumps({"_version": INDEX_VERSION, "documents": index}),
            encoding="utf-8"
        )
    total = sum(len(d["chunks"]) for d in index.values())
    return f"{len(index)} documents, {total} chunks ({changed} updated)"


def search(query: str, k: int = 4) -> list[tuple[float, str, str]]:
    index = _load()
    items = [(name, c) for name, d in index.items() for c in d["chunks"]]
    if not items:
        return []
    q = embed([query], "search_query: ")[0]
    scored = [(_cos(q, c["vec"]), name, c["text"]) for name, c in items]
    return sorted(scored, reverse=True)[:k]


if __name__ == "__main__":
    print(build_index())